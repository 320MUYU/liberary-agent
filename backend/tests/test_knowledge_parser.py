import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from app.common.exceptions import BusinessException
from app.services import knowledge_parser


class KnowledgeParserTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.directory = Path(self.temp_dir.name)

    def write_text(self, name, content):
        path = self.directory / name
        path.write_text(content, encoding="utf-8")
        return path

    def test_utf8_text_is_split_with_chinese_boundaries(self):
        path = self.write_text("rules.txt", "读者应当保持安静。" * 20)

        chunks = knowledge_parser.parse_document(path, chunk_size=40, chunk_overlap=10)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(0 < len(chunk.text) <= 40 for chunk in chunks))
        self.assertTrue(all(chunk.page is None for chunk in chunks))
        self.assertTrue(all("读者" in chunk.text for chunk in chunks))

    def test_markdown_headings_are_preserved_as_source_sections(self):
        path = self.write_text("rules.MD", "# 馆规\n请保持安静。\n\n## 预约\n每天最多预约两小时。")

        chunks = knowledge_parser.parse_document(path)

        self.assertIn("馆规", chunks[0].section)
        self.assertIn("预约", chunks[-1].section)
        self.assertIn("每天最多预约两小时", chunks[-1].text)

    def test_text_bom_is_supported(self):
        path = self.write_text("rules.txt", "\ufeff请保持安静。")

        chunks = knowledge_parser.parse_document(path)

        self.assertEqual(chunks[0].text, "请保持安静。")

    def test_docx_keeps_paragraph_order_and_table_column_labels(self):
        from docx import Document

        document = Document()
        document.add_heading("预约说明", level=1)
        document.add_paragraph("请提前预约。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "类型"
        table.cell(0, 1).text = "时长"
        table.cell(1, 0).text = "个人预约"
        table.cell(1, 1).text = "两小时"
        document.add_paragraph("使用结束请签退。")
        path = self.directory / "rules.docx"
        document.save(path)

        chunks = knowledge_parser.parse_document(path)
        text = "\n".join(chunk.text for chunk in chunks)

        self.assertLess(text.index("请提前预约"), text.index("个人预约"))
        self.assertLess(text.index("个人预约"), text.index("使用结束"))
        self.assertIn("类型：个人预约", text)
        self.assertIn("时长：两小时", text)
        self.assertTrue(all(chunk.section for chunk in chunks))

    def test_docx_zip_bomb_is_rejected_before_document_loading(self):
        path = self.directory / "oversized.docx"
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "a" * (20 * 1024 * 1024 + 1))

        with self.assertRaises(BusinessException) as raised:
            knowledge_parser.parse_document(path)

        self.assertIn("解压", raised.exception.message)

    def test_docx_with_macro_payload_is_rejected(self):
        path = self.directory / "macro.docx"
        with ZipFile(path, "w") as archive:
            archive.writestr("word/vbaProject.bin", "macro payload")

        with self.assertRaises(BusinessException) as raised:
            knowledge_parser.parse_document(path)

        self.assertIn("宏", raised.exception.message)

    def test_pdf_keeps_page_numbers_and_never_merges_pages(self):
        path = self.write_text("rules.pdf", "placeholder")
        reader = SimpleNamespace(
            is_encrypted=False,
            pages=[
                SimpleNamespace(extract_text=lambda: "第一页馆规。" * 8),
                SimpleNamespace(extract_text=lambda: "第二页预约。" * 8),
            ],
        )
        with patch.object(knowledge_parser, "PdfReader", return_value=reader):
            chunks = knowledge_parser.parse_document(path, chunk_size=30, chunk_overlap=5)

        self.assertEqual({chunk.page for chunk in chunks}, {1, 2})
        for chunk in chunks:
            self.assertIn("第一页" if chunk.page == 1 else "第二页", chunk.text)
            self.assertNotIn("第二页" if chunk.page == 1 else "第一页", chunk.text)

    def test_scanned_and_encrypted_pdf_have_actionable_errors(self):
        path = self.write_text("rules.pdf", "placeholder")
        cases = [
            (SimpleNamespace(is_encrypted=True), "加密"),
            (
                SimpleNamespace(
                    is_encrypted=False,
                    pages=[SimpleNamespace(extract_text=lambda: "")],
                ),
                "扫描",
            ),
        ]
        for reader, expected_message in cases:
            with self.subTest(expected_message=expected_message):
                with patch.object(knowledge_parser, "PdfReader", return_value=reader):
                    with self.assertRaises(BusinessException) as raised:
                        knowledge_parser.parse_document(path)
                self.assertIn(expected_message, raised.exception.message)

    def test_pdf_page_limit_is_checked_before_extracting(self):
        path = self.write_text("rules.pdf", "placeholder")
        reader = SimpleNamespace(is_encrypted=False, pages=[object(), object()])
        with patch.object(knowledge_parser, "PdfReader", return_value=reader):
            with self.assertRaises(BusinessException) as raised:
                knowledge_parser.parse_document(path, max_pages=1)

        self.assertIn("页数", raised.exception.message)

    def test_character_limit_applies_across_pdf_pages(self):
        path = self.write_text("rules.pdf", "placeholder")
        reader = SimpleNamespace(
            is_encrypted=False,
            pages=[
                SimpleNamespace(extract_text=lambda: "馆规" * 10),
                SimpleNamespace(extract_text=lambda: "预约" * 10),
            ],
        )
        with patch.object(knowledge_parser, "PdfReader", return_value=reader):
            with self.assertRaises(BusinessException) as raised:
                knowledge_parser.parse_document(path, max_chars=30)

        self.assertIn("字数", raised.exception.message)

    def test_empty_text_and_excessive_characters_are_rejected(self):
        for content, max_chars, message in [(" \n\t", 100, "内容为空"), ("馆规" * 10, 10, "字数")]:
            with self.subTest(message=message):
                path = self.write_text("rules.txt", content)
                with self.assertRaises(BusinessException) as raised:
                    knowledge_parser.parse_document(path, max_chars=max_chars)
                self.assertIn(message, raised.exception.message)

    def test_unsupported_types_and_invalid_utf8_are_rejected_without_path(self):
        for name in ["rules.doc", "rules.docm", "rules.exe", "rules.txt"]:
            with self.subTest(name=name):
                path = self.directory / name
                path.write_bytes(b"\xff\xfe\x00")
                with self.assertRaises(BusinessException) as raised:
                    knowledge_parser.parse_document(path)
                self.assertNotIn(str(path), raised.exception.message)

    def test_missing_and_malformed_documents_hide_internal_errors(self):
        for name in ["missing.txt", "broken.pdf", "broken.docx"]:
            with self.subTest(name=name):
                path = self.directory / name
                if name != "missing.txt":
                    path.write_bytes(b"not a document")
                with self.assertRaises(BusinessException) as raised:
                    knowledge_parser.parse_document(path)
                self.assertNotIn(str(path), raised.exception.message)

    def test_invalid_chunk_configuration_is_rejected(self):
        path = self.write_text("rules.txt", "请保持安静。")
        for options in [
            {"chunk_size": 0},
            {"chunk_size": 20, "chunk_overlap": 20},
            {"chunk_overlap": -1},
            {"max_chars": 0},
            {"max_pages": 0},
        ]:
            with self.subTest(options=options):
                with self.assertRaises(BusinessException):
                    knowledge_parser.parse_document(path, **options)


if __name__ == "__main__":
    unittest.main()
